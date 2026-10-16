# -*- coding: utf-8 -*-
"""正文审计 · 清单构建：把正文里每一处数字断句、每一张表的每个单元格、
每一条 Section/Table/Figure 指针都抽出来，落成文件供逐条回算。"""
import sys, re, os
sys.stdout.reconfigure(encoding='utf-8')
import docx

HERE = r'D:/ml/paper/draft/round6'
OUT = r'D:/ml/chem_ml5/_audit'
os.makedirs(OUT, exist_ok=True)
d = docx.Document(os.path.join(HERE, '终稿_ACS_r6.docx'))
P = [p.text for p in d.paragraphs]

NUM = re.compile(r'\d')
# ---------- 1. 含数字/指针的断句 ----------
sent_re = re.compile(r'(?<=[.;:!?])\s+(?=[A-Z(\u0394])')
with open(os.path.join(OUT, 'sent.txt'), 'w', encoding='utf-8') as f:
    n = 0
    for i, t in enumerate(P):
        if not t.strip():
            continue
        # 断句（保留分隔符）
        parts = re.split(r'(?<=[.;])\s+', t)
        for s in parts:
            if NUM.search(s) or re.search(r'Section \d|Table S?\d|Figure \d', s):
                f.write('P%03d| %s\n' % (i, s.strip()))
                n += 1
    print('含数字/指针断句 %d 条 -> _audit/sent.txt' % n)

# ---------- 2. 表 ----------
with open(os.path.join(OUT, 'tables.txt'), 'w', encoding='utf-8') as f:
    for ti, tb in enumerate(d.tables):
        f.write('\n' + '=' * 90 + '\n表 %d  (%d x %d)\n' % (ti + 1, len(tb.rows), len(tb.columns)))
        for r in tb.rows:
            f.write(' | '.join(c.text.replace('\n', ' ') for c in r.cells) + '\n')
    print('表 %d 张 -> _audit/tables.txt' % len(d.tables))

# ---------- 3. 指针 ----------
ptr = re.compile(r'(Section[s]? [^),;.]{1,30}?|Table S?\d[^),;.]{0,20}|Figure \d[^),;.]{0,20}|'
                 r'Supporting Information[^),;.]{0,40})')
with open(os.path.join(OUT, 'ptr.txt'), 'w', encoding='utf-8') as f:
    for i, t in enumerate(P):
        for m in ptr.finditer(t):
            a, b = max(0, m.start() - 55), min(len(t), m.end() + 40)
            f.write('P%03d| %-34s … %s …\n' % (i, m.group(1).strip(),
                                               t[a:b].replace('\n', ' ')))
    print('指针清单 -> _audit/ptr.txt')

# ---------- 4. 段落长度表（找异常/短段） ----------
with open(os.path.join(OUT, 'paras.txt'), 'w', encoding='utf-8') as f:
    for i, t in enumerate(P):
        f.write('P%03d %5d  %s\n' % (i, len(t), t[:90].replace('\n', ' ')))
print('段落表 -> _audit/paras.txt')
